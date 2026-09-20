from flask import (
    Flask,
    render_template,
    request,
    redirect,
    session,
    send_from_directory
)

from db import Base, engine, SessionLocal
import models

import PyPDF2
import docx
import json
import os

from werkzeug.utils import secure_filename


# =====================================================
# FLASK APP
# =====================================================

app = Flask(__name__)

app.secret_key = "secret123"


# =====================================================
# UPLOAD FOLDER
# =====================================================

UPLOAD_FOLDER = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "uploads"
)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)


# =====================================================
# DATABASE
# =====================================================

Base.metadata.create_all(
    bind=engine
)


# =====================================================
# CSS FILE
# style.css is inside templates folder
# =====================================================

@app.route("/style.css")
def style_css():

    css_path = os.path.join(
        app.template_folder,
        "style.css"
    )

    if not os.path.exists(css_path):
        return (
            "style.css not found in templates folder",
            404
        )

    return send_from_directory(
        app.template_folder,
        "style.css",
        mimetype="text/css"
    )


# =====================================================
# HOME
# =====================================================

@app.route("/")
def home():

    # Logged-in user -> Dashboard
    if "user_id" in session:
        return redirect("/dashboard")

    # Logged-out user -> Landing Page
    return render_template(
        "index.html"
    )


# =====================================================
# SIGNUP
# =====================================================

@app.route(
    "/signup",
    methods=["GET", "POST"]
)
def signup():

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        if not username or not email or not password:

            return """
            <h2>Please fill all fields.</h2>
            <a href="/signup">Back to Signup</a>
            """

        db = SessionLocal()

        try:

            existing_user = (
                db.query(models.User)
                .filter_by(email=email)
                .first()
            )

            if existing_user:

                return """
                <h2>User already exists.</h2>
                <p>Please login with your existing account.</p>
                <a href="/login">Go to Login</a>
                """

            new_user = models.User(
                username=username,
                email=email,
                password=password
            )

            db.add(new_user)

            db.commit()

            return redirect("/login")

        except Exception as e:

            db.rollback()

            print(
                "SIGNUP ERROR:",
                repr(e)
            )

            return f"""
            <h2>Signup Error</h2>
            <p>{e}</p>
            <a href="/signup">Back to Signup</a>
            """

        finally:

            db.close()

    return render_template(
        "signup.html"
    )


# =====================================================
# LOGIN
# =====================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if request.method == "GET":

        try:

            return render_template(
                "login.html"
            )

        except Exception as e:

            print(
                "LOGIN PAGE ERROR:",
                repr(e)
            )

            return f"""
            <h2>Login Page Error</h2>
            <p>{e}</p>
            """


    email = request.form.get(
        "email",
        ""
    ).strip()

    password = request.form.get(
        "password",
        ""
    )

    if not email or not password:

        return """
        <h2>Please enter email and password.</h2>
        <a href="/login">Back to Login</a>
        """


    db = SessionLocal()

    try:

        user = (
            db.query(models.User)
            .filter_by(email=email)
            .first()
        )

        if user and user.password == password:

            session["user_id"] = user.id

            session["user_email"] = user.email

            session["username"] = user.username

            session.modified = True

            return redirect("/dashboard")


        return """
        <h2>Invalid email or password.</h2>
        <a href="/login">Try Again</a>
        """


    except Exception as e:

        print(
            "LOGIN ERROR:",
            repr(e)
        )

        return f"""
        <h2>Login Error</h2>
        <p>{e}</p>
        """


    finally:

        db.close()


# =====================================================
# PDF TEXT EXTRACTION
# =====================================================

def extract_pdf_text(file):

    text = ""

    try:

        reader = PyPDF2.PdfReader(
            file
        )

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text += (
                    page_text +
                    "\n"
                )

    except Exception as e:

        print(
            "PDF ERROR:",
            repr(e)
        )

    return text


# =====================================================
# DOCX TEXT EXTRACTION
# =====================================================

def extract_docx_text(file):

    text = ""

    try:

        document = docx.Document(
            file
        )

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text += (
                    paragraph.text +
                    "\n"
                )

    except Exception as e:

        print(
            "DOCX ERROR:",
            repr(e)
        )

    return text


# =====================================================
# RESUME ANALYSIS
# =====================================================

def analyze_resume(
    resume_text,
    role
):

    resume_lower = resume_text.lower()

    role_lower = role.lower()


    # =================================================
    # ROLE SKILLS
    # =================================================

    role_skills = {

        "data scientist": [

            "python",
            "sql",
            "statistics",
            "machine learning",
            "pandas",
            "numpy",
            "scikit-learn",
            "data visualization",
            "matplotlib",
            "seaborn",
            "tensorflow",
            "power bi",
            "tableau"

        ],


        "data analyst": [

            "python",
            "sql",
            "excel",
            "statistics",
            "pandas",
            "numpy",
            "power bi",
            "tableau",
            "data visualization",
            "matplotlib"

        ],


        "machine learning engineer": [

            "python",
            "machine learning",
            "deep learning",
            "tensorflow",
            "pytorch",
            "scikit-learn",
            "numpy",
            "pandas",
            "sql",
            "docker",
            "git"

        ],


        "software engineer": [

            "java",
            "python",
            "c++",
            "javascript",
            "html",
            "css",
            "sql",
            "git",
            "react",
            "node"

        ],


        "web developer": [

            "html",
            "css",
            "javascript",
            "bootstrap",
            "react",
            "node",
            "sql",
            "git"

        ]

    }


    # =================================================
    # SELECT ROLE
    # =================================================

    selected_skills = role_skills.get(
        role_lower,
        role_skills["data scientist"]
    )


    # =================================================
    # FIND SKILLS
    # =================================================

    found_skills = []

    missing_skills = []


    for skill in selected_skills:

        if skill.lower() in resume_lower:

            found_skills.append(
                skill
            )

        else:

            missing_skills.append(
                skill
            )


    # =================================================
    # SCORE
    # =================================================

    if selected_skills:

        score = int(
            (
                len(found_skills)
                /
                len(selected_skills)
            ) * 100
        )

    else:

        score = 0


    # =================================================
    # ROADMAP
    # =================================================

    roadmap = []


    if missing_skills:

        roadmap.append(
            "Learn the missing technical skills."
        )

        roadmap.append(
            "Build 2-3 practical projects related to the target role."
        )

        roadmap.append(
            "Practice aptitude and technical interview questions."
        )

        roadmap.append(
            "Improve your resume with measurable project achievements."
        )

    else:

        roadmap.append(
            "Your resume covers the major skills for this role."
        )

        roadmap.append(
            "Focus on advanced projects and interview preparation."
        )

        roadmap.append(
            "Practice coding, SQL and technical interview questions."
        )


    # =================================================
    # GENERAL INTERVIEW QUESTIONS
    # =================================================

    interview_questions = [

        f"Tell me about yourself and your experience relevant to {role}.",

        f"Why are you interested in the {role} position?",

        "Explain one important project mentioned in your resume.",

        "What technical skills are you strongest in?",

        "Describe a difficult problem you faced in a project and how you solved it.",

        "How do you keep yourself updated with new technologies?",

        "What are your strengths and areas you are currently improving?"

    ]


    # =================================================
    # DATA SCIENTIST QUESTIONS
    # =================================================

    if "data scientist" in role_lower:

        interview_questions.extend([

            "What is the difference between supervised and unsupervised learning?",

            "Explain the bias-variance tradeoff.",

            "What is overfitting and how can you prevent it?",

            "Explain precision, recall and F1-score.",

            "What is the purpose of cross-validation?",

            "Explain the difference between classification and regression.",

            "How does feature scaling affect machine learning models?"

        ])


    # =================================================
    # DATA ANALYST QUESTIONS
    # =================================================

    elif "data analyst" in role_lower:

        interview_questions.extend([

            "What is the difference between mean, median and mode?",

            "How do you handle missing data?",

            "What is the difference between WHERE and HAVING in SQL?",

            "Explain INNER JOIN and LEFT JOIN.",

            "How do you identify trends in a dataset?",

            "How would you create an effective dashboard?",

            "What is the purpose of data visualization?"

        ])


    # =================================================
    # MACHINE LEARNING ENGINEER QUESTIONS
    # =================================================

    elif "machine learning engineer" in role_lower:

        interview_questions.extend([

            "What is the difference between machine learning and deep learning?",

            "Explain overfitting and regularization.",

            "What is cross-validation?",

            "Explain gradient descent.",

            "What is the difference between TensorFlow and PyTorch?",

            "How would you deploy a machine learning model?",

            "Why is Docker useful for machine learning applications?"

        ])


    # =================================================
    # SOFTWARE ENGINEER QUESTIONS
    # =================================================

    elif "software engineer" in role_lower:

        interview_questions.extend([

            "Explain object-oriented programming concepts.",

            "What is the difference between an array and a linked list?",

            "Explain stack and queue.",

            "What is time complexity?",

            "Explain inheritance and polymorphism.",

            "What is the difference between SQL and NoSQL?",

            "Explain Git and version control."

        ])


    # =================================================
    # WEB DEVELOPER QUESTIONS
    # =================================================

    elif "web developer" in role_lower:

        interview_questions.extend([

            "What is the difference between HTML, CSS and JavaScript?",

            "What is responsive web design?",

            "What is the DOM?",

            "What is React?",

            "What is the difference between frontend and backend?",

            "What is an API?",

            "How do you improve website performance?"

        ])


    # =================================================
    # FINAL RESULT
    # =================================================

    result = {

        "found_skills": found_skills,

        "missing_skills": missing_skills,

        "roadmap": roadmap,

        "interview_questions": interview_questions,

        "score": score,

        "role": role

    }


    return result


# =====================================================
# SAVE RESUME TEXT
# =====================================================

def save_resume_for_user(
    resume_text
):

    user_id = session["user_id"]

    filename = (
        f"user_{user_id}_resume.txt"
    )

    filepath = os.path.abspath(
        os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )
    )

    with open(
        filepath,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            resume_text
        )


    session["resume_text_file"] = filename

    session["resume_available"] = True

    session.modified = True


# =====================================================
# CREATE CORRECTED RESUME
# =====================================================

def create_corrected_resume(
    resume_text
):

    corrected_lines = []


    # =================================================
    # COMMON RESUME HEADINGS
    # =================================================

    heading_words = {

        "summary",
        "professional summary",
        "career objective",
        "objective",
        "education",
        "skills",
        "technical skills",
        "projects",
        "experience",
        "work experience",
        "internship",
        "internships",
        "certifications",
        "achievements",
        "languages",
        "contact",
        "profile",
        "about me"

    }


    # =================================================
    # CLEAN RESUME TEXT
    # =================================================

    for raw_line in resume_text.splitlines():

        line = raw_line.strip()


        # Skip empty lines

        if not line:
            continue


        # Remove extra spaces

        line = " ".join(
            line.split()
        )


        # Normalize bullet points

        if line.startswith("*"):

            line = (
                "• " +
                line.lstrip("* ")
            )

        elif line.startswith("-"):

            line = (
                "• " +
                line.lstrip("- ")
            )


        # Detect headings

        clean_heading = (
            line
            .lower()
            .rstrip(":")
        )


        if clean_heading in heading_words:

            corrected_lines.append(
                (
                    "HEADING",
                    line.rstrip(":")
                )
            )

        else:

            corrected_lines.append(
                (
                    "TEXT",
                    line
                )
            )


    # =================================================
    # CREATE DOCX
    # =================================================

    document = docx.Document()


    # =================================================
    # PAGE MARGINS
    # =================================================

    section = document.sections[0]

    section.top_margin = (
        docx.shared.Inches(0.6)
    )

    section.bottom_margin = (
        docx.shared.Inches(0.6)
    )

    section.left_margin = (
        docx.shared.Inches(0.7)
    )

    section.right_margin = (
        docx.shared.Inches(0.7)
    )


    # =================================================
    # TITLE
    # =================================================

    title = document.add_heading(
        "Improved Resume",
        level=0
    )

    title.alignment = 1


    # =================================================
    # ADD CONTENT
    # =================================================

    for content_type, content in corrected_lines:

        if content_type == "HEADING":

            document.add_heading(
                content,
                level=1
            )

        else:

            paragraph = document.add_paragraph(
                content
            )

            paragraph.paragraph_format.space_after = (
                docx.shared.Pt(5)
            )


    # =================================================
    # SAVE DOCX
    # =================================================

    filename = (
        f"Corrected_Resume_"
        f"{session['user_id']}.docx"
    )

    filepath = os.path.abspath(
        os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )
    )


    document.save(
        filepath
    )


    # =================================================
    # VERIFY FILE WAS CREATED
    # =================================================

    if not os.path.isfile(filepath):

        raise FileNotFoundError(
            "Corrected resume file was not created."
        )


    print(
        "CORRECTED RESUME CREATED:",
        filepath
    )


    return filename


# =====================================================
# DASHBOARD
# =====================================================

@app.route(
    "/dashboard",
    methods=["GET", "POST"]
)
def dashboard():

    if "user_id" not in session:
        return redirect("/login")


    result = None


    # =================================================
    # POST - ANALYZE
    # =================================================

    if request.method == "POST":

        resume_text = request.form.get(
            "resume",
            ""
        ).strip()


        role = request.form.get(
            "role",
            ""
        ).strip()


        uploaded_file = request.files.get(
            "resume_file"
        )


        # =================================================
        # TARGET ROLE REQUIRED
        # =================================================

        if not role:

            result = {

                "error":
                "Please enter your target role."

            }

            return render_template(
                "dashboard.html",
                user=session.get("user_email"),
                username=session.get("username"),
                result=result
            )


        # =================================================
        # FILE UPLOAD
        # =================================================

        if (
            uploaded_file
            and
            uploaded_file.filename
        ):

            original_filename = secure_filename(
                uploaded_file.filename
            )


            filename_lower = (
                original_filename.lower()
            )


            try:

                # =================================================
                # PDF
                # =================================================

                if filename_lower.endswith(
                    ".pdf"
                ):

                    resume_text = extract_pdf_text(
                        uploaded_file
                    )


                # =================================================
                # DOCX
                # =================================================

                elif filename_lower.endswith(
                    ".docx"
                ):

                    resume_text = extract_docx_text(
                        uploaded_file
                    )


                # =================================================
                # OLD DOC
                # =================================================

                elif filename_lower.endswith(
                    ".doc"
                ):

                    result = {

                        "error":
                        "Old .doc files are not supported. "
                        "Please upload .pdf or .docx."

                    }

                    return render_template(
                        "dashboard.html",
                        user=session.get("user_email"),
                        username=session.get("username"),
                        result=result
                    )


                # =================================================
                # INVALID FILE
                # =================================================

                else:

                    result = {

                        "error":
                        "Please upload a PDF or DOCX file."

                    }

                    return render_template(
                        "dashboard.html",
                        user=session.get("user_email"),
                        username=session.get("username"),
                        result=result
                    )


            except Exception as e:

                print(
                    "FILE ERROR:",
                    repr(e)
                )


                result = {

                    "error":
                    f"Unable to read the uploaded file: {e}"

                }


                return render_template(
                    "dashboard.html",
                    user=session.get("user_email"),
                    username=session.get("username"),
                    result=result
                )


        # =================================================
        # CHECK RESUME
        # =================================================

        if not resume_text:

            result = {

                "error":
                "Please paste your resume or upload a PDF/DOCX file."

            }


            return render_template(
                "dashboard.html",
                user=session.get("user_email"),
                username=session.get("username"),
                result=result
            )


        # =================================================
        # SAVE RESUME
        # =================================================

        try:

            save_resume_for_user(
                resume_text
            )

        except Exception as e:

            print(
                "RESUME SAVE ERROR:",
                repr(e)
            )


        # =================================================
        # REMOVE OLD CORRECTED RESUME
        # =================================================

        old_corrected = session.get(
            "corrected_resume"
        )


        if old_corrected:

            old_path = os.path.abspath(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    os.path.basename(
                        old_corrected
                    )
                )
            )


            try:

                if os.path.isfile(
                    old_path
                ):

                    os.remove(
                        old_path
                    )

                    print(
                        "OLD CORRECTED RESUME REMOVED:",
                        old_path
                    )

            except Exception as e:

                print(
                    "OLD FILE REMOVE ERROR:",
                    repr(e)
                )


        session.pop(
            "corrected_resume",
            None
        )

        session.modified = True


        # =================================================
        # ANALYZE RESUME
        # =================================================

        result = analyze_resume(
            resume_text,
            role
        )


        # =================================================
        # SAVE REPORT TO DATABASE
        # =================================================

        db = SessionLocal()


        try:

            report = models.Report(

                user_id=session["user_id"],

                report_text=resume_text,

                result=json.dumps(
                    result
                )

            )


            db.add(
                report
            )

            db.commit()


        except Exception as e:

            db.rollback()

            print(
                "REPORT SAVE ERROR:",
                repr(e)
            )


        finally:

            db.close()


    # =================================================
    # DASHBOARD PAGE
    # =================================================

    return render_template(
        "dashboard.html",
        user=session.get("user_email"),
        username=session.get("username"),
        result=result
    )


# =====================================================
# HISTORY
# =====================================================

@app.route("/history")
def history():

    if "user_id" not in session:
        return redirect("/login")


    db = SessionLocal()


    try:

        reports = (
            db.query(models.Report)
            .filter(
                models.Report.user_id
                ==
                session["user_id"]
            )
            .order_by(
                models.Report.id.desc()
            )
            .all()
        )


        return render_template(
            "history.html",
            reports=reports
        )


    except Exception as e:

        print(
            "HISTORY ERROR:",
            repr(e)
        )


        return f"""
        <h2>History Error</h2>

        <p>{e}</p>

        <a href="/dashboard">
            Back to Dashboard
        </a>
        """


    finally:

        db.close()


# =====================================================
# FIX & IMPROVE RESUME
# =====================================================

@app.route("/fix-resume")
def fix_resume():

    if "user_id" not in session:
        return redirect("/login")


    try:

        # =================================================
        # GET SAVED RESUME
        # =================================================

        filename = session.get(
            "resume_text_file"
        )


        if not filename:

            return """
            <h2>No Resume Found</h2>

            <p>
                Please analyze your resume first.
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # SECURITY
        # =================================================

        filename = os.path.basename(
            filename
        )


        # =================================================
        # BUILD FILE PATH
        # =================================================

        filepath = os.path.abspath(
            os.path.join(
                app.config["UPLOAD_FOLDER"],
                filename
            )
        )


        print(
            "ORIGINAL RESUME PATH:",
            filepath
        )


        # =================================================
        # CHECK FILE
        # =================================================

        if not os.path.isfile(
            filepath
        ):

            return f"""
            <h2>Resume File Not Found</h2>

            <p>
                {filepath}
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # READ RESUME
        # =================================================

        with open(
            filepath,
            "r",
            encoding="utf-8"
        ) as file:

            resume_text = file.read()


        if not resume_text.strip():

            return """
            <h2>Resume is Empty</h2>

            <p>
                Please upload or paste your resume again.
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # CREATE CORRECTED RESUME
        # =================================================

        corrected_filename = (
            create_corrected_resume(
                resume_text
            )
        )


        # =================================================
        # VERIFY CORRECTED RESUME
        # =================================================

        corrected_filename = os.path.basename(
            corrected_filename
        )


        corrected_path = os.path.abspath(
            os.path.join(
                app.config["UPLOAD_FOLDER"],
                corrected_filename
            )
        )


        print(
            "CORRECTED RESUME PATH:",
            corrected_path
        )


        if not os.path.isfile(
            corrected_path
        ):

            return """
            <h2>Correction Failed</h2>

            <p>
                Corrected resume could not be created.
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # SAVE IN SESSION
        # =================================================

        session["corrected_resume"] = (
            corrected_filename
        )

        session.modified = True


        print(
            "SESSION CORRECTED RESUME:",
            session.get(
                "corrected_resume"
            )
        )


        # =================================================
        # RETURN TO DASHBOARD
        # =================================================

        return redirect(
            "/dashboard"
        )


    except Exception as e:

        print(
            "FIX RESUME ERROR:",
            repr(e)
        )


        return f"""
        <h2>Resume Correction Error</h2>

        <p>{e}</p>

        <br>

        <a href="/dashboard">
            Back to Dashboard
        </a>
        """


# =====================================================
# DOWNLOAD CORRECTED RESUME
# =====================================================

@app.route(
    "/download-corrected-resume"
)
def download_corrected_resume():

    if "user_id" not in session:
        return redirect("/login")


    try:

        # =================================================
        # GET CORRECTED FILE
        # =================================================

        filename = session.get(
            "corrected_resume"
        )


        print(
            "DOWNLOAD REQUEST:",
            filename
        )


        # =================================================
        # CHECK SESSION
        # =================================================

        if not filename:

            return """
            <h2>No Corrected Resume Available</h2>

            <p>
                Please click
                <strong>
                    Fix & Improve Resume
                </strong>
                first.
            </p>

            <br>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # SECURITY
        # =================================================

        filename = os.path.basename(
            filename
        )


        # Only DOCX files allowed

        if not filename.lower().endswith(
            ".docx"
        ):

            return """
            <h2>Invalid Resume File</h2>

            <p>
                Only DOCX corrected resumes
                can be downloaded.
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # BUILD ABSOLUTE PATH
        # =================================================

        filepath = os.path.abspath(
            os.path.join(
                app.config["UPLOAD_FOLDER"],
                filename
            )
        )


        print(
            "DOWNLOAD FILE PATH:",
            filepath
        )


        # =================================================
        # CHECK FILE EXISTS
        # =================================================

        if not os.path.isfile(
            filepath
        ):

            return f"""
            <h2>Corrected Resume Not Found</h2>

            <p>
                The corrected resume file
                does not exist.
            </p>

            <p>
                File:
                {filename}
            </p>

            <p>
                Path:
                {filepath}
            </p>

            <br>

            <a href="/dashboard">
                Back to Dashboard
            </a>
            """


        # =================================================
        # DOWNLOAD FILE
        # =================================================

        return send_from_directory(

            directory=os.path.dirname(
                filepath
            ),

            path=os.path.basename(
                filepath
            ),

            as_attachment=True,

            download_name=(
                "Corrected_Resume.docx"
            ),

            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )


    except Exception as e:

        print(
            "DOWNLOAD ERROR:",
            repr(e)
        )


        return f"""
        <h2>Download Error</h2>

        <p>
            {e}
        </p>

        <br>

        <a href="/dashboard">
            Back to Dashboard
        </a>
        """


# =====================================================
# LOGOUT
# =====================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect("/login")


# =====================================================
# RUN APPLICATION
# =====================================================

if __name__ == "__main__":

    app.run(

        host="127.0.0.1",

        port=5000,

        debug=True,

        use_reloader=False
    )